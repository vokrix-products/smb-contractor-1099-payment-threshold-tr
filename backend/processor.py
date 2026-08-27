import csv
import io
import json
import os
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import pdfplumber
from openpyxl import load_workbook
from openpyxl import Workbook
from openai import OpenAI

STATUS_VALUES = {
    "w9_missing:critical",
    "w9_valid:good",
    "w9_expired_stale:warning",
    "w9_flagged:critical",
    "w9_awaiting_signature:warning",
    "w9_not_verified:warning",
    "tin_mismatch:critical",
    "payment_threshold_approaching:warning",
    "payment_threshold_crossed:info",
    "over_threshold_without_w9:critical",
    "backup_withholding_risk:critical",
    "chase_active:info",
    "chase_overdue:warning",
    "chase_failed:critical",
    "contractor_duplicate:warning",
    "contractor_incomplete_data:warning",
    "non_us_payee:info",
    "archived_paid_under_threshold:good",
    "ready_to_file:good",
}

ALIAS_MAP = {
    "contractor_name": [
        "contractor", "payee", "payee_name", "contractor_name", "name",
        "vendor", "vendor_name", "supplier", "recipient", "recipient_name",
        "business_name", "legal_name", "customer", "client", "company_name",
        "operator", "provider", "merchant", "seller",
    ],
    "legal_business_name": [
        "legal_business_name", "dba", "doing_business_as", "business_name",
        "company", "trade_name",
    ],
    "ein_ssn": [
        "ein", "ssn", "tax_id", "tin", "employer_identification_number",
        "social_security_number", "taxpayer_id", "federal_id", "taxid",
        "tax_identification_number",
    ],
    "w9_entity_type": [
        "w9_entity_type", "entity_type", "business_type", "w_9_entity_type",
        "legal_entity_type", "tax_classification",
    ],
    "w9_status": ["w9_status", "w_9_status", "status"],
    "w9_received_date": [
        "w9_received_date", "w_9_received_date", "received_date", "date_received",
    ],
    "w9_signature_date": [
        "w9_signature_date", "signature_date", "signed_date", "date_signed",
    ],
    "w9_form_version_exempt": [
        "w9_form_version", "form_version", "exempt_status", "w9_version",
        "w_9_form_version",
    ],
    "contractor_email": [
        "contractor_email", "email", "email_address", "e_mail", "payee_email",
    ],
    "contractor_phone": [
        "contractor_phone", "phone", "phone_number", "mobile", "telephone",
    ],
    "contractor_address": [
        "contractor_address", "mailing_address", "address", "street_address",
        "city_state_zip", "location",
    ],
    "tin_match_status": [
        "tin_match_status", "tin_match", "irs_match", "taxpayer_match",
    ],
    "payment_date": [
        "payment_date", "pay_date", "date", "transaction_date", "paid_date",
        "posted_date", "payment_posted_date", "date_paid",
    ],
    "payment_amount": [
        "payment_amount", "amount", "paid", "total", "payment", "gross_amount",
        "net_amount", "price", "amount_paid", "transaction_amount",
    ],
    "payment_source_platform": [
        "payment_source", "source", "platform", "payment_source_platform",
        "payment_platform", "pay_source",
    ],
    "payment_method": ["payment_method", "method", "pay_method"],
    "transaction_id": [
        "transaction_id", "transaction", "reference", "reference_number",
        "txn_id", "payment_reference", "confirmation_id",
    ],
    "invoice_number_description": [
        "invoice_number", "invoice", "description", "memo", "notes",
        "reference_description",
    ],
    "calendar_year": ["calendar_year", "year", "payment_year", "tax_year"],
    "cumulative_ytd": [
        "cumulative_ytd", "ytd_total", "cumulative_payment", "ytd",
        "total_ytd", "year_to_date",
    ],
    "threshold_status": ["threshold_status", "threshold"],
    "form_1099_type": [
        "form_1099_type", "1099_form_type", "form_type", "1099_type",
    ],
    "backup_withholding": [
        "backup_withholding", "backup_withholding_flag",
        "backup_withholding_rate", "withholding",
    ],
    "last_w9_request_date": [
        "last_w9_request_date", "last_request_date", "w9_last_request_date",
    ],
    "w9_chase_step": ["w9_chase_step", "chase_step", "chase_status"],
    "next_chase_due_date": [
        "next_chase_due_date", "chase_due_date", "next_due_date",
    ],
    "risk_status": ["risk_status", "risk"],
}


def _normalize_header(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^0-9a-z_]+", "_", value)
    return value.strip("_")


def _lookup_key(norm: str) -> str:
    if norm in ALIAS_MAP:
        return norm
    for canonical, aliases in ALIAS_MAP.items():
        if norm in aliases:
            return canonical
    return norm


def _stringify(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _parse_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    if not raw:
        return None

    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%d %b %Y", "%b %d, %Y", "%m/%d/%y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass

    try:
        year = int(raw)
        if 1900 <= year <= 2100:
            return date(year, 12, 31)
    except ValueError:
        pass

    return None


def _parse_amount(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    raw = str(value).replace("$", "").replace(",", "").replace(" ", "")
    if not raw:
        return 0.0

    negative = False
    if raw.startswith("(") and raw.endswith(")"):
        negative = True
        raw = raw[1:-1]

    try:
        amount = float(raw)
        return -amount if negative else amount
    except ValueError:
        match = re.search(r"[-+]?\\d[\\d,]*\\.?\\d*", raw)
        if match:
            try:
                amount = float(match.group().replace(",", ""))
                return -amount if negative else amount
            except ValueError:
                return 0.0
    return 0.0


def _lower(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _normalize_record(raw: Dict[str, Any]) -> Dict[str, Any]:
    record: Dict[str, Any] = {}
    for key, value in raw.items():
        if value is None:
            continue
        norm = _normalize_header(str(key))
        canonical = _lookup_key(norm)
        val = _stringify(value).strip()
        if not val:
            continue
        if canonical in record and str(record[canonical]) != val:
            record[canonical] = str(record[canonical]) + " | " + val
        else:
            record[canonical] = val

    title = (
        record.get("contractor_name")
        or record.get("legal_business_name")
        or record.get("ein_ssn")
        or "Unknown Payee"
    )
    record["title"] = str(title).strip()
    return record


def _normalize_all(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    normalized: List[Dict[str, Any]] = []
    for row in rows:
        record = _normalize_record(row)
        if record:
            normalized.append(record)
    return normalized


def _is_w9_record(record: Dict[str, Any]) -> bool:
    w9_fields = [
        "w9_status",
        "w9_received_date",
        "w9_signature_date",
        "w9_entity_type",
        "w9_form_version_exempt",
    ]
    if any(record.get(field) for field in w9_fields):
        return True
    if (
        (record.get("ein_ssn") or record.get("contractor_email") or record.get("contractor_address"))
        and not record.get("payment_amount")
    ):
        return True
    return False


def _derive_status(record: Dict[str, Any]) -> str:
    w9_status = _lower(record.get("w9_status"))
    tin_match = _lower(record.get("tin_match_status"))
    risk = _lower(record.get("risk_status"))
    chase_step = _lower(record.get("w9_chase_step"))
    backup = record.get("backup_withholding")
    threshold = _lower(record.get("threshold_status"))
    form_type = _lower(record.get("form_1099_type"))

    cumulative_raw = record.get("cumulative_ytd") or record.get("payment_amount")
    cumulative = _parse_amount(cumulative_raw)

    if _is_w9_record(record):
        if tin_match == "mismatch":
            return "tin_mismatch:critical"
        if w9_status in {"flagged"} or "flag" in risk:
            return "w9_flagged:critical"
        if w9_status in {"expired", "stale"} or "expired" in w9_status:
            return "w9_expired_stale:warning"
        if not record.get("w9_signature_date") and (
            w9_status in {"awaiting", "pending", "awaiting signature", "not signed"}
            or "awaiting" in risk
        ):
            return "w9_awaiting_signature:warning"
        if w9_status == "valid" or record.get("w9_received_date"):
            if tin_match == "valid" or not record.get("tin_match_status"):
                return "w9_valid:good"
            return "w9_not_verified:warning"
        return "w9_not_verified:warning"

    if backup and cumulative >= 600:
        return "backup_withholding_risk:critical"
    if tin_match == "mismatch":
        return "tin_mismatch:critical"
    if w9_status in {"flagged"}:
        return "w9_flagged:critical"

    has_w9 = bool(
        record.get("w9_status")
        or record.get("w9_received_date")
        or record.get("w9_signature_date")
    )

    if not has_w9:
        if cumulative >= 600:
            return "over_threshold_without_w9:critical"
        return "w9_missing:critical"

    if chase_step and chase_step != "none":
        if "fail" in chase_step or "fail" in risk:
            return "chase_failed:critical"
        next_chase = _parse_date(record.get("next_chase_due_date"))
        if next_chase and next_chase < date.today():
            return "chase_overdue:warning"
        return "chase_active:info"

    if threshold:
        if "cross" in threshold:
            return "payment_threshold_crossed:info"
        if "approach" in threshold:
            return "payment_threshold_approaching:warning"
        if "not" in threshold or "under" in threshold:
            return "archived_paid_under_threshold:good"

    if cumulative >= 600:
        if w9_status == "valid" or record.get("w9_received_date"):
            if form_type in {"1099-nec", "1099-misc", "1099 nec", "1099 misc"}:
                return "ready_to_file:good"
            return "payment_threshold_crossed:info"
        return "over_threshold_without_w9:critical"

    if cumulative >= 500:
        return "payment_threshold_approaching:warning"

    return "archived_paid_under_threshold:good"


def _due_date_for_record(record: Dict[str, Any]) -> Optional[str]:
    for field in (
        "next_chase_due_date",
        "w9_signature_date",
        "w9_received_date",
        "payment_date",
        "last_w9_request_date",
    ):
        parsed = _parse_date(record.get(field))
        if parsed:
            return parsed.isoformat()
    return None


def _finalize_record(record: Dict[str, Any]) -> Dict[str, Any]:
    out = {
        "title": record.get("title", "Unknown Payee"),
        "status": record.get("status", "contractor_incomplete_data:warning"),
        "due_date": record.get("due_date"),
        "details": {},
    }

    for key, value in record.items():
        if key in {"title", "status", "due_date"}:
            continue
        if value is None:
            continue
        if isinstance(value, datetime):
            value = value.isoformat()
        elif isinstance(value, date):
            value = value.isoformat()
        elif isinstance(value, bytes):
            value = value.decode("utf-8", errors="ignore")
        out["details"][key] = value

    return out


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_text = []
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
            return "\n".join(pages_text).strip()
    except Exception:
        return ""


def _extract_xlsx(file_bytes: bytes) -> List[Dict[str, Any]]:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return []

    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return []

    if not rows:
        return []

    headers = []
    for value in rows[0]:
        headers.append(_normalize_header(str(value)) if value is not None else "")
    if not any(headers):
        headers = [f"column_{i}" for i in range(len(rows[0]))]

    output: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if all(value is None for value in raw):
            continue
        row: Dict[str, Any] = {}
        for index, value in enumerate(raw):
            header = headers[index] if index < len(headers) else f"column_{index}"
            if value is not None:
                row[header] = _stringify(value)
        output.append(row)

    return output


def _process_text(text: str) -> List[Dict[str, Any]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return []

    sample = "\n".join(lines[:10])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t|;")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        rows: List[Dict[str, Any]] = []
        for row in reader:
            if not any(value for value in row.values()):
                continue
            clean = {
                str(key): (str(value).strip() if value is not None else "")
                for key, value in row.items()
            }
            rows.append(clean)
        if rows and any(key for key in rows[0].keys() if key):
            return rows
    except Exception:
        pass

    record: Dict[str, str] = {}
    for line in lines:
        if ":" in line:
            key, value = line.split(":", 1)
            key = key.strip()
            value = value.strip()
            if key and value:
                record[key] = value
            continue
        if "," in line:
            parts = [part.strip() for part in line.split(",")]
            if len(parts) >= 3:
                if not record:
                    record["contractor_name"] = parts[0]
                    record["payment_date"] = parts[1]
                    record["payment_amount"] = parts[2]
                    if len(parts) > 3:
                        record["payment_source_platform"] = parts[3]
            continue
        if "description" not in record:
            record["description"] = line

    return [record] if record else []


def process_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    if not file_bytes:
        return []

    text = ""
    if file_bytes[:4] == b"%PDF":
        text = _extract_pdf(file_bytes)
        if not text.strip():
            return []
        rows = _process_text(text)
    else:
        try:
            rows = _extract_xlsx(file_bytes)
        except Exception:
            rows = []
        if not rows:
            try:
                text = file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                text = ""
            rows = _process_text(text)

    if not rows:
        return []

    normalized = _normalize_all(rows)

    records: List[Dict[str, Any]] = []
    for record in normalized:
        status = _derive_status(record)
        due_date = _due_date_for_record(record)
        record["status"] = status
        record["due_date"] = due_date
        records.append(_finalize_record(record))

    return records

def extract_text(file_bytes):
    try:
        import pdfplumber, io
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            text = ""; [text := text + (p.extract_text() or "") + "\n" for p in pdf.pages]
            if text.strip(): return text
    except: pass
    return file_bytes.decode("utf-8", errors="ignore")
